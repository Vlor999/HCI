"""Turns the files of the documents folder (Markdown, text, PDF, DOCX, Pages) into retrievable chunks.

Structure first, then size: a document is split into sections (Markdown and DOCX headings, PDF pages), each section
into paragraphs, then the paragraphs' sentences are packed into chunks of at most `chunk_words` words, with
`chunk_overlap_words` words of overlap between neighbours of the same section. A chunk never spans two sections, and
a sentence is only cut when it is longer than a chunk on its own.
"""

import shutil
import subprocess
import tempfile
import unicodedata
from collections.abc import Callable, Iterable
from pathlib import Path

import docx
from loguru import logger
from pypdf import PdfReader

from hci.config import get_settings
from hci.constants import (
    DOCUMENT_EXTENSIONS,
    MARKDOWN_HEADING,
    PAGES_EXPORT_SCRIPT,
    PARAGRAPH_BREAK,
    PATH_SEPARATORS,
    ROOT_DIR,
    SENTENCE_END,
    UNSAFE_NAME_CHARS,
    UPLOADS_DIRNAME,
    Source,
)
from hci.types import DocumentSection, RecordChunk


def file_ref(path: Path) -> str:
    """Reference to a file for chunk ids: its path relative to the repo root, or absolute outside the repo.

    Args:
        path: The file.

    Returns:
        The POSIX path, e.g. `data/paths.json`.
    """
    path = path.resolve()
    return path.relative_to(ROOT_DIR).as_posix() if path.is_relative_to(ROOT_DIR) else path.as_posix()


def document_files() -> list[Path]:
    """Supported files of the documents folder, recursively, hidden files and folders left out.

    Returns:
        The files, sorted by path.
    """
    root = get_settings().documents_dir
    return sorted(
        p
        for p in root.rglob("*")
        if p.is_file()
        and p.suffix.lower() in DOCUMENT_EXTENSIONS
        and not any(part.startswith(".") for part in p.relative_to(root).parts)
    )


def safe_name(filename: str) -> str:
    """A client file name made safe to save: no folders, no leading dots, only word characters, `.` and `-`.

    Args:
        filename: The name sent by the client, possibly with a path (`../../x.md`, `C:\\x.md`).

    Returns:
        The base name, possibly empty.
    """
    base = PATH_SEPARATORS.split(filename)[-1]
    return UNSAFE_NAME_CHARS.sub("_", base).lstrip(".")


def save_upload(name: str, data: bytes) -> Path:
    """Save an uploaded document in the uploads subfolder, never overwriting: `name-1.ext`, `name-2.ext`...

    Args:
        name: A name from `safe_name`.
        data: The file content.

    Returns:
        The saved file.
    """
    folder = get_settings().documents_dir / UPLOADS_DIRNAME
    folder.mkdir(parents=True, exist_ok=True)
    stem, suffix = Path(name).stem, Path(name).suffix
    n = 0
    while True:
        path = folder / (f"{stem}-{n}{suffix}" if n else name)
        try:
            with path.open("xb") as f:  # exclusive create: two uploads of the same name never overwrite each other
                f.write(data)
            return path
        except FileExistsError:
            n += 1


def load_documents() -> list[RecordChunk]:
    """Chunks of every supported file of the documents folder.

    Returns:
        The chunks, file by file in path order.
    """
    return [chunk for path in document_files() for chunk in load_document(path)]


def load_document(path: Path) -> list[RecordChunk]:
    """Chunks of one document, ids `<file ref>#<n>` counted from 1. A file that fails to parse gives no chunks.

    Args:
        path: The document, with a supported extension.

    Returns:
        The chunks, each starting with `<file name>, <section or page>: `.
    """
    try:
        sections = PARSERS[path.suffix.lower()](path)
    # pypdf and python-docx raise many types on malformed files; one bad file never stops the index
    except Exception as e:
        logger.warning("Skipping document {}: cannot parse it ({})", path, e)
        return []
    settings = get_settings()
    bodies = [
        (s.title, body)
        for s in sections
        for body in pack(s.paragraphs, settings.chunk_words, settings.chunk_overlap_words)
    ]
    ref = file_ref(path)
    return [
        RecordChunk(f"{ref}#{n}", f"{path.name}{', ' + title if title else ''}: {body}", Source.DOCUMENT)
        for n, (title, body) in enumerate(bodies, 1)
    ]


def pack(paragraphs: list[str], limit: int, overlap: int) -> list[str]:
    """Pack paragraphs into chunks of at most `limit` words, cutting between sentences.

    Args:
        paragraphs: The section's paragraphs.
        limit: Max words per chunk.
        overlap: Max words of whole trailing sentences repeated at the start of the next chunk.

    Returns:
        The chunk bodies.
    """
    units = [piece for p in paragraphs for s in SENTENCE_END.split(p) for piece in _windows(s, limit)]
    chunks: list[list[str]] = []
    current: list[str] = []
    for unit in units:
        if current and _count([*current, unit]) > limit:
            chunks.append(current)
            current = _tail(current, overlap)
            if _count([*current, unit]) > limit:
                current = []
        current.append(unit)
    return [" ".join(c) for c in [*chunks, current] if c]


def _count(units: Iterable[str]) -> int:
    return sum(len(u.split()) for u in units)


def _windows(sentence: str, limit: int) -> list[str]:
    """A sentence as is, or cut every `limit` words when it is longer than a chunk on its own."""
    words = sentence.split()
    return [" ".join(words[i : i + limit]) for i in range(0, len(words), limit)]


def _tail(units: list[str], overlap: int) -> list[str]:
    """The trailing units that fit in `overlap` words."""
    tail: list[str] = []
    for unit in reversed(units):
        if _count([unit, *tail]) > overlap:
            break
        tail.insert(0, unit)
    return tail


def _paragraphs(text: str) -> list[str]:
    """Blank-line paragraphs, whitespace collapsed; NFKC turns PDF ligatures (`ﬂ`) back into letters for BM25."""
    text = unicodedata.normalize("NFKC", text)
    return [" ".join(p.split()) for p in PARAGRAPH_BREAK.split(text) if p.strip()]


def _sections(blocks: Iterable[tuple[int, str]]) -> list[DocumentSection]:
    """Group `(heading level, text)` blocks into sections titled by their heading path; level 0 is body text."""
    sections, path, body = [], [], []
    for level, text in blocks:
        if not level:
            body.append(text)
            continue
        sections.append(DocumentSection(" > ".join(path), _paragraphs("\n".join(body))))
        path, body = [*path[: level - 1], text.strip()], []
    sections.append(DocumentSection(" > ".join(path), _paragraphs("\n".join(body))))
    return [s for s in sections if s.paragraphs]


def _markdown_blocks(text: str) -> Iterable[tuple[int, str]]:
    # ponytail: a `#` line inside a code fence reads as a heading; track fences if code-heavy docs show up
    for line in text.splitlines():
        m = MARKDOWN_HEADING.match(line)
        yield (len(m[1]), m[2]) if m else (0, line)


def _docx_level(style_name: str) -> int:
    """Heading level of a DOCX paragraph style: `Title` 1, `Heading <n>` n, anything else 0 (body)."""
    if style_name == "Title":
        return 1
    level = style_name.removeprefix("Heading ")
    return int(level) if level != style_name and level.isdigit() else 0


def _markdown_sections(path: Path) -> list[DocumentSection]:
    return _sections(_markdown_blocks(path.read_text(errors="replace")))


def _text_sections(path: Path) -> list[DocumentSection]:
    return [DocumentSection("", paragraphs)] if (paragraphs := _paragraphs(path.read_text(errors="replace"))) else []


def _docx_sections(path: Path) -> list[DocumentSection]:
    paragraphs = docx.Document(str(path)).paragraphs
    # a blank line after each body paragraph keeps DOCX paragraphs apart in `_paragraphs`
    return _sections((_docx_level(p.style.name if p.style else ""), p.text + "\n") for p in paragraphs)


def _pdf_sections(path: Path) -> list[DocumentSection]:
    pages = PdfReader(path).pages
    sections = [DocumentSection(f"page {n}", _paragraphs(page.extract_text() or "")) for n, page in enumerate(pages, 1)]
    return [s for s in sections if s.paragraphs]


def _pages_sections(path: Path) -> list[DocumentSection]:
    with tempfile.TemporaryDirectory() as tmp:
        pdf = export_pages(path, Path(tmp) / f"{path.stem}.pdf")
        return _pdf_sections(pdf) if pdf else []


def export_pages(source: Path, target: Path) -> Path | None:
    """Export an Apple Pages file to PDF through the Pages app (macOS only).

    Args:
        source: The `.pages` file.
        target: The PDF to write.

    Returns:
        `target`, or None (with a warning) when Pages is not available or the export failed.
    """
    hint = "export it from Pages as PDF or DOCX and add that file instead"
    if not shutil.which("osascript"):
        logger.warning("Skipping {}: .pages files need the Pages app on macOS; {}", source, hint)
        return None
    command = ["osascript", "-e", PAGES_EXPORT_SCRIPT, str(source.resolve()), str(target)]
    try:
        subprocess.run(command, check=True, capture_output=True, timeout=get_settings().pages_export_timeout)
    except (OSError, subprocess.SubprocessError) as e:
        logger.warning("Skipping {}: the Pages app could not export it ({}); {}", source, e, hint)
        return None
    return target


PARSERS: dict[str, Callable[[Path], list[DocumentSection]]] = {
    ".md": _markdown_sections,
    ".txt": _text_sections,
    ".pdf": _pdf_sections,
    ".docx": _docx_sections,
    ".pages": _pages_sections,
}
